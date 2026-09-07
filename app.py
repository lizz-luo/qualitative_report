import streamlit as st
import pandas as pd
import io
import os
import tempfile
import pyreadstat
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

st.set_page_config(page_title="學校問卷回應整理工具（SPSS）", layout="wide")
st.title("學校問卷回應整理工具（SPSS .sav）")

st.markdown("""
本工具支援兩種 SPSS 檔案：

**Q1Q2 模式（Q1Q2.sav）**  
- 輸入欄位：SchName, Q2_12, Q1_12, Q1Q2_13, Q1Q2_14, Q1Q2_15  
- 輸出表格：SchName x Q1Q2_12, SchName x Q1Q2_13, SchName x Q1Q2_14, SchName x Q1Q2_15  

**Q3 模式（Q3.sav）**  
- 輸入欄位：SchName, Q3_23, Q3_25, Q3_26  
- 輸出表格：SchName x Q3_23, SchName x Q3_25, SchName x Q3_26  

每張表會保留每位老師的一條回應；同一學校只在首次出現時顯示名稱，並在 Word 中合併後續空白單元格。
""")

INVALID_VALUES = {"", "nil", "/", "nan", "none", "n/a", "na"}

COLUMNS_Q1Q2 = ["SchName", "Q2_12", "Q1_12", "Q1Q2_13", "Q1Q2_14", "Q1Q2_15"]
COLUMNS_Q3 = ["SchName", "Q3_23", "Q3_25", "Q3_26"]


def is_valid_response(value):
    """判斷回應是否有有效文字。"""
    if pd.isna(value):
        return False
    return str(value).strip().lower() not in INVALID_VALUES


def combine_responses(value1, value2):
    """合併兩個問題欄的有效文字。"""
    parts = []
    for value in (value1, value2):
        if is_valid_response(value):
            parts.append(str(value).strip())
    return " ".join(parts)


def make_question_table(data, question_column):
    """建立指定問題的兩欄表格，並讓連續重複的學校名稱只顯示一次。"""
    output = data[["SchName", question_column]].copy()
    output = output[output[question_column].apply(is_valid_response)].copy()
    output = output.reset_index(drop=True)

    output["學校"] = output["SchName"].where(
        output["SchName"].ne(output["SchName"].shift()), ""
    )
    output = output[["學校", question_column]]
    return output


def read_sav_file(uploaded_file):
    """把 Streamlit 上傳檔暫存後讀為 DataFrame。"""
    file_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".sav") as temp_file:
            temp_file.write(uploaded_file.getvalue())
            file_path = temp_file.name

        raw_df, metadata = pyreadstat.read_sav(file_path)
        return raw_df, metadata
    finally:
        if file_path and os.path.exists(file_path):
            os.remove(file_path)


def create_word_document(tables_dict_of_dict):
    """
    建立 Word 文件，包含多個模式的表格。
    tables_dict_of_dict: {mode_name: {question_name: df}}
    """
    doc = Document()

    title = doc.add_heading("學校問卷回應整理報告", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for mode_name, tables in tables_dict_of_dict.items():
        doc.add_heading(mode_name, level=2)

        for question_name, df in tables.items():
            doc.add_heading(question_name, level=3)

            table = doc.add_table(rows=len(df) + 1, cols=2)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            table.columns[0].width = Inches(2.5)
            table.columns[1].width = Inches(5.0)

            header_cells = table.rows[0].cells
            header_cells[0].text = "學校"
            header_cells[1].text = question_name

            for cell in header_cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.size = Pt(10)

            for row_index, (_, row_data) in enumerate(df.iterrows(), start=1):
                row = table.rows[row_index]
                row.cells[0].text = str(row_data["學校"]) if row_data["學校"] else ""
                row.cells[1].text = str(row_data[question_name])

                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        for run in paragraph.runs:
                            run.font.size = Pt(9)

            school_values = df["學校"].tolist()
            merge_ranges = []
            current_start = None

            for i, school in enumerate(school_values, start=1):
                if school:
                    if current_start is not None and current_start + 1 < i:
                        merge_ranges.append((current_start, i - 1))
                    current_start = i

            if current_start is not None and current_start < len(school_values) + 1:
                merge_ranges.append((current_start, len(school_values)))

            for start_row, end_row in merge_ranges:
                if end_row > start_row:
                    first_cell = table.cell(start_row, 0)
                    for row_idx in range(start_row + 1, end_row + 1):
                        cell_to_merge = table.cell(row_idx, 0)
                        first_cell.merge(cell_to_merge)

            doc.add_paragraph()

    return doc


tab_main, tab_q3 = st.tabs(["Q1Q2 模式", "Q3 模式"])

# ==================== Q1Q2 模式 ====================
with tab_main:
    st.subheader("Q1Q2 模式（Q1Q2.sav）")
    uploaded_q1q2 = st.file_uploader("上傳 Q1Q2.sav 檔案", type=["sav"], key="q1q2_uploader")

    if uploaded_q1q2 is not None:
        try:
            raw_df, metadata = read_sav_file(uploaded_q1q2)

            if not isinstance(raw_df, pd.DataFrame):
                st.error(f"無法把 .sav 檔讀取成資料表，取得的類型：{type(raw_df)}")
                st.stop()

            st.subheader("原始資料預覽")
            st.caption(f"共讀取 {len(raw_df):,} 筆資料、{len(raw_df.columns):,} 個變量。")
            st.dataframe(raw_df.head(20), use_container_width=True)

            missing_columns = [column for column in COLUMNS_Q1Q2 if column not in raw_df.columns]
            if missing_columns:
                st.error("檔案缺少必要變量：" + ", ".join(missing_columns))
                st.write("目前讀取到的變量：", list(raw_df.columns))
                st.stop()

            data = raw_df[COLUMNS_Q1Q2].copy()
            data["Q1Q2_12"] = data.apply(
                lambda row: combine_responses(row["Q2_12"], row["Q1_12"]), axis=1
            )

            question_columns = ["Q1Q2_12", "Q1Q2_13", "Q1Q2_14", "Q1Q2_15"]
            tables = {
                question: make_question_table(data, question)
                for question in question_columns
            }

            st.subheader("四張整理後表格")
            tabs = st.tabs(question_columns)

            for tab, question in zip(tabs, question_columns):
                with tab:
                    result_df = tables[question]
                    st.caption(f"有效回應：{len(result_df):,} 條")
                    st.dataframe(result_df, use_container_width=True, height=500)

                    csv_data = result_df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")
                    st.download_button(
                        label=f"下載 {question}.csv",
                        data=csv_data,
                        file_name=f"SchName_x_{question}.csv",
                        mime="text/csv",
                        key=f"download_q1q2_{question}",
                    )

            excel_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
                for question, result_df in tables.items():
                    result_df.to_excel(writer, index=False, sheet_name=question)
            excel_buffer.seek(0)

            st.download_button(
                label="一次下載四張表（Excel）",
                data=excel_buffer,
                file_name="SchName_Q1Q2_tables.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="download_q1q2_excel",
            )

            doc_buffer = io.BytesIO()
            doc = create_word_document({"Q1Q2 模式": tables})
            doc.save(doc_buffer)
            doc_buffer.seek(0)

            st.download_button(
                label="下載整理後的 Word 報告（Q1Q2）",
                data=doc_buffer,
                file_name="SchName_Q1Q2_report.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="download_q1q2_word",
            )

        except Exception as error:
            st.error(f"讀取或處理檔案時發生錯誤：{error}")
            st.exception(error)
    else:
        st.info("請上傳包含 SchName, Q2_12, Q1_12, Q1Q2_13, Q1Q2_14, Q1Q2_15 的 .sav 檔案。")

# ==================== Q3 模式 ====================
with tab_q3:
    st.subheader("Q3 模式（Q3.sav）")
    uploaded_q3 = st.file_uploader("上傳 Q3.sav 檔案", type=["sav"], key="q3_uploader")

    if uploaded_q3 is not None:
        try:
            raw_df, metadata = read_sav_file(uploaded_q3)

            if not isinstance(raw_df, pd.DataFrame):
                st.error(f"無法把 .sav 檔讀取成資料表，取得的類型：{type(raw_df)}")
                st.stop()

            st.subheader("原始資料預覽")
            st.caption(f"共讀取 {len(raw_df):,} 筆資料、{len(raw_df.columns):,} 個變量。")
            st.dataframe(raw_df.head(20), use_container_width=True)

            missing_columns = [column for column in COLUMNS_Q3 if column not in raw_df.columns]
            if missing_columns:
                st.error("檔案缺少必要變量：" + ", ".join(missing_columns))
                st.write("目前讀取到的變量：", list(raw_df.columns))
                st.stop()

            data = raw_df[COLUMNS_Q3].copy()

            question_columns = ["Q3_23", "Q3_25", "Q3_26"]
            tables = {
                question: make_question_table(data, question)
                for question in question_columns
            }

            st.subheader("三張整理後表格")
            tabs = st.tabs(question_columns)

            for tab, question in zip(tabs, question_columns):
                with tab:
                    result_df = tables[question]
                    st.caption(f"有效回應：{len(result_df):,} 條")
                    st.dataframe(result_df, use_container_width=True, height=500)

                    csv_data = result_df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")
                    st.download_button(
                        label=f"下載 {question}.csv",
                        data=csv_data,
                        file_name=f"SchName_x_{question}.csv",
                        mime="text/csv",
                        key=f"download_q3_{question}",
                    )

            excel_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
                for question, result_df in tables.items():
                    result_df.to_excel(writer, index=False, sheet_name=question)
            excel_buffer.seek(0)

            st.download_button(
                label="一次下載三張表（Excel）",
                data=excel_buffer,
                file_name="SchName_Q3_tables.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="download_q3_excel",
            )

            doc_buffer = io.BytesIO()
            doc = create_word_document({"Q3 模式": tables})
            doc.save(doc_buffer)
            doc_buffer.seek(0)

            st.download_button(
                label="下載整理後的 Word 報告（Q3）",
                data=doc_buffer,
                file_name="SchName_Q3_report.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="download_q3_word",
            )

        except Exception as error:
            st.error(f"讀取或處理檔案時發生錯誤：{error}")
            st.exception(error)
    else:
        st.info("請上傳包含 SchName, Q3_23, Q3_25, Q3_26 的 .sav 檔案。")